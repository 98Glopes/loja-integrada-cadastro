"""Prova de conceito: gera uma planilha de importação da Loja Integrada com dados dummy.

Uso:
    python poc/gerar_planilha_poc.py [--exportacao docs/produtos-....xlsx] [--saida poc/saida]

O layout de colunas (nomes e ordem) é lido da linha 1 da exportação real da loja,
porque só ela contém as grades cadastradas na plataforma (`grade-tamanho-infantil`,
`grade-produto-com-uma-cor`). Os produtos gerados têm SKU com prefixo `POC-` e
`ativo = S` para serem conferidos no site e apagados depois.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import openpyxl

EXPORTACAO_PADRAO = Path("docs/brutos/produtos-2026-09-10-03-25-78670adce0a94f6.xlsx")
SAIDA_PADRAO = Path("poc/saida")
TOTAL_COLUNAS_ESPERADO = 54

# URLs públicas das imagens por SKU do produto pai (até 5), hospedadas no registry.
BASE_URL_IMAGENS = "https://bucket.gleite.com/fotos"
IMAGENS: dict[str, list[str]] = {
    "POC-VEST-001": [
        f"{BASE_URL_IMAGENS}/somnii-vestido-florzinha-beige-frente.jpg",
        f"{BASE_URL_IMAGENS}/somnii-vestido-florzinha-rosa-frente.jpg",
    ],
    "POC-CONJ-002": [f"{BASE_URL_IMAGENS}/kiki-conjunto-body-calca-branco-frente.jpg"],
    "POC-BLUS-003": [f"{BASE_URL_IMAGENS}/onda-marinha-blusinha-canelada-preto-frente.jpg"],
    # Rodada 2: isolar por que POC-VEST-001 (2 imagens) ficou sem foto.
    "POC-VEST-004": [f"{BASE_URL_IMAGENS}/somnii-vestido-florzinha-beige-frente.jpg"],
    "POC-VEST-005": [
        f"{BASE_URL_IMAGENS}/kiki-conjunto-body-calca-branco-frente.jpg",
        f"{BASE_URL_IMAGENS}/onda-marinha-blusinha-canelada-preto-frente.jpg",
    ],
}


@dataclass(frozen=True)
class ProdutoDummy:
    sku: str
    nome: str
    marca: str
    seo_tag_title: str
    seo_tag_description: str
    descricao_html: str
    cores: tuple[str, ...]
    tamanhos: tuple[str, ...]
    preco_cheio: float
    estoque_por_variacao: int = 1
    peso_em_kg: float = 0.1
    altura_em_cm: int = 4
    largura_em_cm: int = 22
    comprimento_em_cm: int = 22
    categoria: tuple[str, ...] = field(default_factory=tuple)


def descricao_html(headline: str, abertura: str, itens: list[str], marca: str) -> str:
    lista = "\n".join(f"  <li>{item}</li>" for item in itens)
    return (
        f"<h2>{headline}</h2>\n\n"
        f"<p>{abertura}</p>\n\n"
        f"<ul>\n{lista}\n</ul>\n\n"
        f"<h3>Sobre a {marca}</h3>\n"
        f"<p>Texto dummy de prova de conceito sobre a marca {marca}.</p>\n\n"
        "<h3>Compre na Kmilaa Modas</h3>\n"
        "<p>Na Kmilaa Modas você encontra uma seleção de moda infantil com atendimento "
        "próximo e envio rápido direto de Sorocaba. Parcele em até 3x sem juros.</p>"
    )


PRODUTOS: tuple[ProdutoDummy, ...] = (
    ProdutoDummy(
        sku="POC-VEST-001",
        nome="POC Vestido Somnii Florzinha Babados e Laço 1 ao 3",
        marca="somnii",
        seo_tag_title="POC Vestido Somnii Florzinha Babados 1 ao 3 | Kmilaa Modas",
        seo_tag_description=(
            "Produto de teste (POC). Vestido infantil somnii em malha canelada com "
            "florzinha, babados e laço. Tamanhos 1 ao 3. Parcele em 3x na Kmilaa Modas."
        ),
        descricao_html=descricao_html(
            "Vestido de teste: leveza e alegria em cada detalhe",
            "Produto dummy da prova de conceito de importação. Malha canelada macia, "
            "babados nas mangas e laço na cintura para os dias de festa e passeio.",
            [
                "Malha canelada 100% algodão",
                "Babados nas mangas",
                "Laço na cintura",
                "Disponível nos tamanhos 1 ao 3",
            ],
            "somnii",
        ),
        cores=("Beige", "Rosa"),
        tamanhos=("1", "2", "3"),
        preco_cheio=119.9,
    ),
    ProdutoDummy(
        sku="POC-CONJ-002",
        nome="POC Conjunto kiki Body e Calça Moletom P ao G",
        marca="kiki",
        seo_tag_title="POC Conjunto kiki Body Calça Moletom P ao G | Kmilaa Modas",
        seo_tag_description=(
            "Produto de teste (POC). Conjunto kiki baby com body em malha e calça de "
            "moletom, 100% algodão. Tamanhos P ao G. Parcele em 3x na Kmilaa Modas."
        ),
        descricao_html=descricao_html(
            "Conjunto de teste: aconchego para os primeiros meses",
            "Produto dummy da prova de conceito de importação. Body em malha macia e "
            "calça de moletom que não pesa e aquece nos dias frios.",
            [
                "100% algodão",
                "Botões no ombro para facilitar a troca",
                "Calça com cós elástico",
                "Disponível nos tamanhos P ao G",
            ],
            "kiki",
        ),
        cores=("Branco",),
        tamanhos=("P", "M", "G"),
        preco_cheio=89.9,
    ),
    ProdutoDummy(
        sku="POC-BLUS-003",
        nome="POC Blusinha Onda Marinha Manga Longa Canelada 4",
        marca="Onda Marinha",
        seo_tag_title="POC Blusinha Onda Marinha Manga Longa 4 | Kmilaa Modas",
        seo_tag_description=(
            "Produto de teste (POC). Blusinha Onda Marinha manga longa em malha canelada "
            "com elastano, tamanho 4, em três cores. Parcele em 3x na Kmilaa Modas."
        ),
        descricao_html=descricao_html(
            "Blusinha de teste: básica que combina com tudo",
            "Produto dummy da prova de conceito de importação. Malha canelada com "
            "elastano que acompanha cada movimento sem apertar.",
            [
                "Malha canelada com elastano",
                "Manga longa",
                "Três opções de cor",
                "Tamanho 4",
            ],
            "Onda Marinha",
        ),
        cores=("Preto", "Vermelho", "Rosa"),
        tamanhos=("4",),
        preco_cheio=59.9,
    ),
    ProdutoDummy(
        sku="POC-VEST-004",
        nome="POC Vestido Somnii Teste Imagem Unica 1 ao 3",
        marca="somnii",
        seo_tag_title="POC Vestido Somnii Teste Imagem Unica 1 ao 3",
        seo_tag_description=(
            "Produto de teste (POC) para isolar importação de imagem: uma única URL, "
            "a mesma que falhou na rodada 1."
        ),
        descricao_html=descricao_html(
            "Vestido de teste: uma imagem",
            "Produto dummy da rodada 2. Mesma imagem beige da rodada 1, sozinha.",
            ["Malha canelada 100% algodão", "Disponível nos tamanhos 1 ao 3"],
            "somnii",
        ),
        cores=("Beige",),
        tamanhos=("1", "2", "3"),
        preco_cheio=119.9,
    ),
    ProdutoDummy(
        sku="POC-VEST-005",
        nome="POC Vestido Somnii Teste Duas Imagens 1 ao 3",
        marca="somnii",
        seo_tag_title="POC Vestido Somnii Teste Duas Imagens 1 ao 3",
        seo_tag_description=(
            "Produto de teste (POC) para isolar importação de imagem: duas URLs que "
            "funcionaram sozinhas na rodada 1."
        ),
        descricao_html=descricao_html(
            "Vestido de teste: duas imagens",
            "Produto dummy da rodada 2. Duas imagens que importaram bem na rodada 1.",
            ["Malha canelada 100% algodão", "Disponível nos tamanhos 1 ao 3"],
            "somnii",
        ),
        cores=("Rosa",),
        tamanhos=("1", "2", "3"),
        preco_cheio=119.9,
    ),
)


def ler_cabecalhos(caminho_exportacao: Path) -> list[str]:
    workbook = openpyxl.load_workbook(caminho_exportacao, read_only=True)
    primeira_linha = next(workbook.worksheets[0].iter_rows(min_row=1, max_row=1, values_only=True))
    cabecalhos = [str(celula) for celula in primeira_linha if celula is not None]
    if len(cabecalhos) != TOTAL_COLUNAS_ESPERADO:
        raise ValueError(
            f"Esperava {TOTAL_COLUNAS_ESPERADO} colunas na exportação, achei {len(cabecalhos)}"
        )
    return cabecalhos


def slug(valor: str) -> str:
    return valor.lower().replace(" ", "-")


def linha_pai(produto: ProdutoDummy) -> dict[str, object]:
    linha: dict[str, object] = {
        "tipo": "com-variacao",
        "sku": produto.sku,
        "ativo": "S",
        "usado": "N",
        "destaque": "N",
        "nome": produto.nome,
        "seo-tag-title": produto.seo_tag_title,
        "seo-tag-description": produto.seo_tag_description,
        "descricao-completa": produto.descricao_html,
        "preco-sob-consulta": "N",
        "marca": produto.marca,
    }
    for nivel, nome_categoria in enumerate(produto.categoria, start=1):
        linha[f"categoria-nome-nivel-{nivel}"] = nome_categoria
    for posicao, url in enumerate(IMAGENS.get(produto.sku, [])[:5], start=1):
        linha[f"imagem-{posicao}"] = url
    return linha


def linha_filha(produto: ProdutoDummy, cor: str, tamanho: str) -> dict[str, object]:
    return {
        "tipo": "variacao",
        "sku-pai": produto.sku,
        "sku": f"{produto.sku}-{slug(cor)}-{slug(tamanho)}",
        "ativo": "S",
        "usado": "N",
        "destaque": "N",
        "estoque-gerenciado": "S",
        "estoque-quantidade": produto.estoque_por_variacao,
        "estoque-situacao-em-estoque": "imediata",
        "estoque-situacao-sem-estoque": "indisponivel",
        "preco-sob-consulta": "N",
        "preco-custo": 0,
        "preco-cheio": produto.preco_cheio,
        "preco-promocional": 0,
        "peso-em-kg": produto.peso_em_kg,
        "altura-em-cm": produto.altura_em_cm,
        "largura-em-cm": produto.largura_em_cm,
        "comprimento-em-cm": produto.comprimento_em_cm,
        "grade-produto-com-uma-cor": cor,
        "grade-tamanho-infantil": tamanho,
    }


def linhas_do_produto(produto: ProdutoDummy) -> list[dict[str, object]]:
    filhas = [
        linha_filha(produto, cor, tamanho) for cor in produto.cores for tamanho in produto.tamanhos
    ]
    return [linha_pai(produto), *filhas]


def escrever_planilha(
    cabecalhos: list[str], linhas: list[dict[str, object]], destino: Path
) -> None:
    for linha in linhas:
        desconhecidas = set(linha) - set(cabecalhos)
        if desconhecidas:
            raise ValueError(f"Colunas fora do layout: {sorted(desconhecidas)}")
    workbook = openpyxl.Workbook()
    sheet = workbook.create_sheet("Sheet1", 0)
    workbook.remove(workbook.worksheets[1])
    sheet.append(cabecalhos)
    for linha in linhas:
        sheet.append([linha.get(coluna) for coluna in cabecalhos])
    destino.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destino)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--exportacao", type=Path, default=EXPORTACAO_PADRAO)
    parser.add_argument("--saida", type=Path, default=SAIDA_PADRAO)
    parser.add_argument("--skus", nargs="+", help="Gera só os produtos pai com estes SKUs")
    parser.add_argument("--sufixo", default="", help="Sufixo no nome do arquivo (ex.: rodada-2)")
    args = parser.parse_args()

    produtos = [p for p in PRODUTOS if not args.skus or p.sku in args.skus]
    cabecalhos = ler_cabecalhos(args.exportacao)
    linhas = [linha for produto in produtos for linha in linhas_do_produto(produto)]
    sufixo = f"-{args.sufixo}" if args.sufixo else ""
    destino = args.saida / f"poc-loja-integrada-{date.today():%Y-%m-%d}{sufixo}.xlsx"
    escrever_planilha(cabecalhos, linhas, destino)

    pais = sum(1 for linha in linhas if linha["tipo"] == "com-variacao")
    variacoes = len(linhas) - pais
    print(f"{destino}: {pais} produtos pai, {variacoes} variações, {len(cabecalhos)} colunas")


if __name__ == "__main__":
    main()
