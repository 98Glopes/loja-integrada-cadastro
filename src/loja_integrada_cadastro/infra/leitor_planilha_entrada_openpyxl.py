from __future__ import annotations

from collections.abc import Sequence
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from loja_integrada_cadastro.models.exceptions.erro_planilha_entrada import ErroPlanilhaEntrada
from loja_integrada_cadastro.models.layout_planilha_entrada import (
    CAMPO_PRODUTO_OPCIONAL,
    CAMPOS_PRODUTO_OBRIGATORIOS,
    CAMPOS_VARIACAO_OBRIGATORIOS,
    COLUNAS_PLANILHA_ENTRADA,
)
from loja_integrada_cadastro.models.problema_linha_planilha import ProblemaLinhaPlanilha
from loja_integrada_cadastro.models.produto_entrada import ProdutoEntrada
from loja_integrada_cadastro.models.variacao_entrada import VariacaoEntrada

_CAMPOS_PRODUTO_TODOS: tuple[str, ...] = (*CAMPOS_PRODUTO_OBRIGATORIOS, CAMPO_PRODUTO_OPCIONAL)


class LeitorPlanilhaEntradaOpenpyxl:
    """Lê a planilha .xlsx de entrada com openpyxl e monta ProdutoEntrada.

    Cabeçalho é validado por nome (ordem livre). Problemas de linha (célula obrigatória vazia,
    conflito de campo de produto) são acumulados e reportados juntos num único
    ErroPlanilhaEntrada; problemas de cabeçalho levantam antes de qualquer linha ser lida.
    """

    def ler(self, caminho: Path) -> list[ProdutoEntrada]:
        workbook = self._abrir(caminho)
        try:
            aba = workbook.active
            assert aba is not None, "planilha sem aba ativa"
            indice = self._validar_cabecalho(caminho, aba)
            return self._processar_linhas(caminho, aba, indice)
        finally:
            workbook.close()

    @staticmethod
    def _abrir(caminho: Path) -> Workbook:
        try:
            return load_workbook(caminho, data_only=True, read_only=True)
        except (InvalidFileException, BadZipFile, FileNotFoundError) as erro:
            raise ErroPlanilhaEntrada(
                caminho,
                (ProblemaLinhaPlanilha(1, "arquivo", f"nao foi possivel abrir: {erro}"),),
            ) from erro

    @staticmethod
    def _validar_cabecalho(caminho: Path, aba: Worksheet) -> dict[str, int]:
        primeira_linha = next(aba.iter_rows(min_row=1, max_row=1, values_only=True))
        encontrado = {
            str(valor).strip(): posicao
            for posicao, valor in enumerate(primeira_linha, start=1)
            if valor is not None and str(valor).strip()
        }
        faltando = [coluna for coluna in COLUNAS_PLANILHA_ENTRADA if coluna not in encontrado]
        desconhecidas = [coluna for coluna in encontrado if coluna not in COLUNAS_PLANILHA_ENTRADA]
        if faltando or desconhecidas:
            problemas = [
                ProblemaLinhaPlanilha(1, coluna, "coluna obrigatoria ausente no cabecalho")
                for coluna in faltando
            ] + [
                ProblemaLinhaPlanilha(1, coluna, "coluna desconhecida no cabecalho")
                for coluna in desconhecidas
            ]
            raise ErroPlanilhaEntrada(caminho, problemas)
        return encontrado

    def _processar_linhas(
        self, caminho: Path, aba: Worksheet, indice: dict[str, int]
    ) -> list[ProdutoEntrada]:
        grupos: dict[str, _GrupoEmConstrucao] = {}
        ordem: list[str] = []
        problemas: list[ProblemaLinhaPlanilha] = []

        for numero_linha, linha in enumerate(aba.iter_rows(min_row=2, values_only=True), start=2):
            valores = self._extrair_valores(linha, indice)
            if all(valor is None for valor in valores.values()):
                continue

            sku_pai = valores["sku-pai"]
            if sku_pai is None:
                problemas.append(
                    ProblemaLinhaPlanilha(numero_linha, "sku-pai", "célula obrigatória vazia")
                )
                continue

            grupo_existente = grupos.get(sku_pai)
            if grupo_existente is None:
                grupo = _GrupoEmConstrucao(sku_pai, numero_linha)
                grupos[sku_pai] = grupo
                ordem.append(sku_pai)
                problemas += self._registrar_primeira_linha_produto(numero_linha, grupo, valores)
            else:
                grupo = grupo_existente
                problemas += self._conferir_linha_seguinte_produto(numero_linha, grupo, valores)

            problemas += self._processar_variacao(numero_linha, grupo, valores)

        if problemas:
            raise ErroPlanilhaEntrada(caminho, problemas)
        return [self._finalizar(grupos[sku_pai]) for sku_pai in ordem]

    @staticmethod
    def _extrair_valores(linha: Sequence[Any], indice: dict[str, int]) -> dict[str, str | None]:
        """Lê cada coluna pela posição do cabeçalho.

        `openpyxl` em modo `read_only` devolve, por linha, uma tupla do tamanho da última
        célula que a planilha guarda naquela linha — se as últimas colunas nunca foram
        digitadas (comum quando a planilha vem do Excel/Sheets), a tupla fica mais curta que o
        cabeçalho. Posição fora da tupla é célula vazia, igual a `None` gravado explicitamente.
        """
        return {
            coluna: (
                LeitorPlanilhaEntradaOpenpyxl._texto(linha[posicao - 1])
                if posicao - 1 < len(linha)
                else None
            )
            for coluna, posicao in indice.items()
        }

    @staticmethod
    def _texto(valor: object) -> str | None:
        if valor is None:
            return None
        if isinstance(valor, float) and valor.is_integer():
            valor = int(valor)
        texto = str(valor).strip()
        return texto or None

    @staticmethod
    def _registrar_primeira_linha_produto(
        numero_linha: int, grupo: _GrupoEmConstrucao, valores: dict[str, str | None]
    ) -> list[ProblemaLinhaPlanilha]:
        """Grava os campos de produto da primeira linha do grupo; exige os obrigatórios."""
        problemas: list[ProblemaLinhaPlanilha] = []
        for campo in _CAMPOS_PRODUTO_TODOS:
            valor = valores.get(campo)
            grupo.campos[campo] = valor
            if valor is None and campo in CAMPOS_PRODUTO_OBRIGATORIOS:
                problemas.append(
                    ProblemaLinhaPlanilha(numero_linha, campo, "célula obrigatória vazia")
                )
        return problemas

    @staticmethod
    def _conferir_linha_seguinte_produto(
        numero_linha: int, grupo: _GrupoEmConstrucao, valores: dict[str, str | None]
    ) -> list[ProblemaLinhaPlanilha]:
        """Confere os campos de produto de uma linha seguinte do grupo contra os já lidos."""
        problemas: list[ProblemaLinhaPlanilha] = []
        for campo in _CAMPOS_PRODUTO_TODOS:
            valor = valores.get(campo)
            if valor is None:
                continue
            valor_atual = grupo.campos.get(campo)
            if valor_atual is None:
                grupo.campos[campo] = valor
            elif valor_atual != valor:
                problemas.append(
                    ProblemaLinhaPlanilha(
                        numero_linha,
                        campo,
                        "valor diferente do ja lido para "
                        f"'{grupo.sku_pai}' (linha {grupo.linha_primeira_ocorrencia}): "
                        f"'{valor_atual}' vs '{valor}'",
                    )
                )
        return problemas

    @staticmethod
    def _processar_variacao(
        numero_linha: int, grupo: _GrupoEmConstrucao, valores: dict[str, str | None]
    ) -> list[ProblemaLinhaPlanilha]:
        problemas: list[ProblemaLinhaPlanilha] = []
        for campo in CAMPOS_VARIACAO_OBRIGATORIOS:
            if valores.get(campo) is None:
                problemas.append(
                    ProblemaLinhaPlanilha(numero_linha, campo, "célula obrigatória vazia")
                )
        if problemas:
            return problemas

        cor = valores["cor"]
        tamanho = valores["tamanho"]
        gtin = valores["gtin"]
        texto_preco = valores["preco"]
        texto_estoque = valores["estoque"]
        assert cor is not None
        assert tamanho is not None
        assert gtin is not None
        assert texto_preco is not None
        assert texto_estoque is not None

        preco = LeitorPlanilhaEntradaOpenpyxl._parse_preco(numero_linha, texto_preco, problemas)
        estoque = LeitorPlanilhaEntradaOpenpyxl._parse_estoque(
            numero_linha, texto_estoque, problemas
        )
        if problemas:
            return problemas

        assert preco is not None
        assert estoque is not None
        grupo.variacoes.append(
            VariacaoEntrada(cor=cor, tamanho=tamanho, gtin=gtin, preco=preco, estoque=estoque)
        )
        return problemas

    @staticmethod
    def _parse_preco(
        numero_linha: int, texto: str, problemas: list[ProblemaLinhaPlanilha]
    ) -> Decimal | None:
        try:
            return Decimal(texto.replace(",", "."))
        except InvalidOperation:
            problemas.append(
                ProblemaLinhaPlanilha(numero_linha, "preco", f"valor não decimal: '{texto}'")
            )
            return None

    @staticmethod
    def _parse_estoque(
        numero_linha: int, texto: str, problemas: list[ProblemaLinhaPlanilha]
    ) -> int | None:
        try:
            return int(texto)
        except ValueError:
            problemas.append(
                ProblemaLinhaPlanilha(numero_linha, "estoque", f"valor não inteiro: '{texto}'")
            )
            return None

    @staticmethod
    def _finalizar(grupo: _GrupoEmConstrucao) -> ProdutoEntrada:
        campos = grupo.campos
        marca = campos["marca"]
        nome_fornecedor = campos["nome-fornecedor"]
        tipo_peca = campos["tipo-peca"]
        categoria_bruta = campos["categoria"]
        composicao = campos["composicao"]
        detalhes = campos["detalhes"]
        faixa_tamanho = campos["faixa-tamanho"]
        assert marca is not None
        assert nome_fornecedor is not None
        assert tipo_peca is not None
        assert categoria_bruta is not None
        assert composicao is not None
        assert detalhes is not None
        assert faixa_tamanho is not None

        categoria = tuple(segmento.strip() for segmento in categoria_bruta.split(">"))
        return ProdutoEntrada(
            sku_pai=grupo.sku_pai,
            marca=marca,
            nome_fornecedor=nome_fornecedor,
            tipo_peca=tipo_peca,
            categoria=categoria,
            composicao=composicao,
            detalhes=detalhes,
            colecao=campos.get(CAMPO_PRODUTO_OPCIONAL),
            faixa_tamanho=faixa_tamanho,
            variacoes=tuple(grupo.variacoes),
        )


class _GrupoEmConstrucao:
    """Estado acumulado durante a leitura de um sku-pai. Não é uma entidade de domínio."""

    def __init__(self, sku_pai: str, linha_primeira_ocorrencia: int) -> None:
        self.sku_pai = sku_pai
        self.linha_primeira_ocorrencia = linha_primeira_ocorrencia
        self.campos: dict[str, str | None] = {}
        self.variacoes: list[VariacaoEntrada] = []
