from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from loja_integrada_cadastro.models.dados_mestre import DadosMestre
from loja_integrada_cadastro.models.layout_planilha_entrada import COLUNAS_PLANILHA_ENTRADA

_LINHA_MAXIMA_VALIDACAO = 1000

_COMENTARIOS: dict[str, str] = {
    "sku-pai": "Codigo do fornecedor. Repita em todas as linhas da mesma familia (mesmo produto).",
    "marca": "Grafia da marca; use a lista suspensa.",
    "nome-fornecedor": "Nome do produto como veio do fornecedor.",
    "tipo-peca": "Ex.: Vestido, Conjunto, Macacao, Casaco...",
    "categoria": (
        "Caminho ate 5 niveis separado por '>' (ex.: Linha Kids (6 ao 10) > Menina > Vestido)."
    ),
    "composicao": "Ex.: 100% algodao.",
    "detalhes": "Diferenciais da peca (botoes, bordado, estampa...).",
    "colecao": "Opcional. Ex.: Outono/Inverno 2026.",
    "faixa-tamanho": ("Texto livre com a faixa de tamanhos da familia, ex.: 'P ao G' ou '1 ao 4'."),
    "cor": ("Grafia exata da cor; use a lista suspensa. Precisa ter subpasta em fotos/<sku-pai>/."),
    "tamanho": "Use a lista suspensa (P M G GG XG ou 1 2 3 4 6 8 10 12 14 16 18 20).",
    "gtin": "Codigo de barras (8/12/13/14 digitos). Digite como TEXTO, nao como numero.",
    "preco": "Decimal maior que zero. Aceita virgula ou ponto (ex.: 119,90 ou 119.9).",
    "estoque": "Numero inteiro maior ou igual a zero.",
}

_EXEMPLO_LINHA_1: dict[str, str] = {
    "sku-pai": "3254002",
    "marca": "kiki",
    "nome-fornecedor": "Conjunto Baby Malha e Moletom",
    "tipo-peca": "Conjunto",
    "categoria": "Linha Baby (P ao XG) > Menino > Conjunto",
    "composicao": "100% algodao",
    "detalhes": "Botoes na gola, bordado no bolso",
    "colecao": "Outono/Inverno 2026",
    "faixa-tamanho": "P ao G",
    "cor": "Beige",
    "tamanho": "P",
    "gtin": "7891234567895",
    "preco": "119,90",
    "estoque": "10",
}

_EXEMPLO_LINHA_2: dict[str, str] = {
    "sku-pai": "3254002",
    "cor": "Beige",
    "tamanho": "G",
    "gtin": "7891234567901",
    "preco": "129,90",
    "estoque": "5",
}


class GeradorModeloEntrada:
    """Gera o modelo .xlsx de entrada, com listas suspensas a partir de DadosMestre."""

    def __init__(self, dados_mestre: DadosMestre) -> None:
        self._dados_mestre = dados_mestre

    def gerar(self, destino: Path) -> None:
        destino.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        aba = workbook.active
        assert aba is not None, "workbook novo sem aba ativa"
        aba.title = "Entrada"

        self._escrever_cabecalho(aba)
        self._escrever_exemplo(aba)
        self._formatar_colunas(aba)

        marcas, cores, tamanhos = self._listas_para_validacao()
        self._criar_aba_listas(workbook, marcas, cores, tamanhos)
        self._aplicar_validacoes(aba, len(marcas), len(cores), len(tamanhos))

        workbook.save(destino)

    @staticmethod
    def _escrever_cabecalho(aba: Worksheet) -> None:
        for posicao, coluna in enumerate(COLUNAS_PLANILHA_ENTRADA, start=1):
            celula = aba.cell(row=1, column=posicao, value=coluna)
            comentario = _COMENTARIOS.get(coluna)
            if comentario:
                celula.comment = Comment(comentario, "loja-integrada-cadastro")

    @staticmethod
    def _escrever_exemplo(aba: Worksheet) -> None:
        for linha_numero, exemplo in ((2, _EXEMPLO_LINHA_1), (3, _EXEMPLO_LINHA_2)):
            for posicao, coluna in enumerate(COLUNAS_PLANILHA_ENTRADA, start=1):
                aba.cell(row=linha_numero, column=posicao, value=exemplo.get(coluna, ""))

    @staticmethod
    def _formatar_colunas(aba: Worksheet) -> None:
        for coluna in ("sku-pai", "gtin"):
            posicao = COLUNAS_PLANILHA_ENTRADA.index(coluna) + 1
            letra = get_column_letter(posicao)
            aba.column_dimensions[letra].number_format = "@"
        aba.freeze_panes = "A2"

    def _listas_para_validacao(self) -> tuple[list[str], list[str], list[str]]:
        marcas = sorted(self._dados_mestre.marcas_canonicas)
        cores = sorted(self._dados_mestre.cores)
        tamanhos = sorted(self._dados_mestre.tamanhos, key=_chave_ordenacao_tamanho)
        return marcas, cores, tamanhos

    @staticmethod
    def _criar_aba_listas(
        workbook: Workbook, marcas: list[str], cores: list[str], tamanhos: list[str]
    ) -> None:
        aba = workbook.create_sheet("Listas")
        aba.append(["marca", "cor", "tamanho"])
        maximo = max(len(marcas), len(cores), len(tamanhos))
        for indice in range(maximo):
            aba.append(
                [
                    marcas[indice] if indice < len(marcas) else None,
                    cores[indice] if indice < len(cores) else None,
                    tamanhos[indice] if indice < len(tamanhos) else None,
                ]
            )
        aba.sheet_state = "hidden"

    @staticmethod
    def _aplicar_validacoes(
        aba: Worksheet, total_marcas: int, total_cores: int, total_tamanhos: int
    ) -> None:
        configuracoes = (
            ("marca", total_marcas, "A"),
            ("cor", total_cores, "B"),
            ("tamanho", total_tamanhos, "C"),
        )
        for coluna, total, letra_lista in configuracoes:
            posicao = COLUNAS_PLANILHA_ENTRADA.index(coluna) + 1
            letra_coluna = get_column_letter(posicao)
            validacao = DataValidation(
                type="list",
                formula1=f"=Listas!${letra_lista}$2:${letra_lista}${total + 1}",
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Valor fora da lista",
                error=f"Selecione um valor da lista para '{coluna}'.",
            )
            aba.add_data_validation(validacao)
            validacao.add(f"{letra_coluna}2:{letra_coluna}{_LINHA_MAXIMA_VALIDACAO}")


def _chave_ordenacao_tamanho(tamanho: str) -> tuple[int, int | str]:
    """Ordena tamanhos numericos por valor e letras alfabeticamente, letras primeiro.

    So para exibicao no modelo; DadosMestre nao guarda ordem canonica (decisao da task 04:
    faixa_tamanho deixou de ser calculada, entao nao ha mais necessidade de ordem cruzando
    escalas).
    """
    if tamanho.isdigit():
        return (1, int(tamanho))
    return (0, tamanho)
