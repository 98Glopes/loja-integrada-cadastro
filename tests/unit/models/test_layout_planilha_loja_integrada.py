from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from loja_integrada_cadastro.models.layout_planilha_loja_integrada import (
    COLUNAS_PLANILHA_SAIDA,
    COLUNAS_PREENCHIDAS_KMILAA,
)

_BRUTOS_DIR = Path(__file__).parents[3] / "docs" / "brutos"


def test_tem_54_colunas() -> None:
    assert len(COLUNAS_PLANILHA_SAIDA) == 54


def test_colunas_sem_duplicatas() -> None:
    assert len(set(COLUNAS_PLANILHA_SAIDA)) == len(COLUNAS_PLANILHA_SAIDA)


def test_colunas_preenchidas_sao_subconjunto_do_layout() -> None:
    assert COLUNAS_PREENCHIDAS_KMILAA <= set(COLUNAS_PLANILHA_SAIDA)


def test_colunas_batem_com_exportacao_real_quando_disponivel() -> None:
    arquivos = sorted(_BRUTOS_DIR.glob("produtos-*.xlsx")) if _BRUTOS_DIR.exists() else []
    if not arquivos:
        pytest.skip("exportação real (docs/brutos/produtos-*.xlsx) não disponível localmente")

    workbook = load_workbook(arquivos[0], read_only=True)
    aba = workbook.active
    assert aba is not None
    primeira_linha = next(aba.iter_rows(min_row=1, max_row=1, values_only=True))

    assert tuple(primeira_linha) == COLUNAS_PLANILHA_SAIDA
