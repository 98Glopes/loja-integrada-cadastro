from pathlib import Path

from openpyxl import load_workbook

from loja_integrada_cadastro.infra.carregador_recursos import CarregadorRecursos
from loja_integrada_cadastro.infra.gerador_modelo_entrada_openpyxl import GeradorModeloEntrada
from loja_integrada_cadastro.infra.leitor_planilha_entrada_openpyxl import (
    LeitorPlanilhaEntradaOpenpyxl,
)
from loja_integrada_cadastro.models.layout_planilha_entrada import COLUNAS_PLANILHA_ENTRADA


def _gerar(tmp_path: Path) -> Path:
    dados_mestre = CarregadorRecursos().dados_mestre()
    destino = tmp_path / "modelo-entrada.xlsx"
    GeradorModeloEntrada(dados_mestre).gerar(destino)
    return destino


def test_cabecalho_tem_as_14_colunas_na_ordem_certa(tmp_path: Path) -> None:
    destino = _gerar(tmp_path)

    workbook = load_workbook(destino)
    aba = workbook["Entrada"]
    cabecalho = tuple(celula.value for celula in next(aba.iter_rows(min_row=1, max_row=1)))

    assert cabecalho == COLUNAS_PLANILHA_ENTRADA


def test_gera_duas_linhas_de_exemplo_do_mesmo_produto(tmp_path: Path) -> None:
    destino = _gerar(tmp_path)

    workbook = load_workbook(destino)
    aba = workbook["Entrada"]
    indice_sku_pai = COLUNAS_PLANILHA_ENTRADA.index("sku-pai")
    linha_2 = [c.value for c in next(aba.iter_rows(min_row=2, max_row=2))]
    linha_3 = [c.value for c in next(aba.iter_rows(min_row=3, max_row=3))]

    assert linha_2[indice_sku_pai] == linha_3[indice_sku_pai]
    assert linha_2[indice_sku_pai]


def test_aba_listas_existe_oculta_com_dados_do_dados_mestre(tmp_path: Path) -> None:
    dados_mestre = CarregadorRecursos().dados_mestre()
    destino = _gerar(tmp_path)

    workbook = load_workbook(destino)
    aba_listas = workbook["Listas"]

    assert aba_listas.sheet_state == "hidden"
    marcas_lidas = {
        linha[0].value for linha in aba_listas.iter_rows(min_row=2) if linha[0].value is not None
    }
    cores_lidas = {
        linha[1].value for linha in aba_listas.iter_rows(min_row=2) if linha[1].value is not None
    }
    assert marcas_lidas == set(dados_mestre.marcas_canonicas)
    assert cores_lidas == set(dados_mestre.cores)


def test_colunas_marca_cor_tamanho_tem_validacao_de_dados(tmp_path: Path) -> None:
    destino = _gerar(tmp_path)

    workbook = load_workbook(destino)
    aba = workbook["Entrada"]
    validacoes = list(aba.data_validations.dataValidation)

    assert len(validacoes) == 3
    for validacao in validacoes:
        assert validacao.type == "list"
        assert "Listas!" in (validacao.formula1 or "")


def test_sku_pai_e_gtin_ficam_formatados_como_texto(tmp_path: Path) -> None:
    destino = _gerar(tmp_path)

    workbook = load_workbook(destino)
    aba = workbook["Entrada"]

    for coluna in ("sku-pai", "gtin"):
        posicao = COLUNAS_PLANILHA_ENTRADA.index(coluna)
        letra = aba.cell(row=1, column=posicao + 1).column_letter
        assert aba.column_dimensions[letra].number_format == "@"


def test_round_trip_gerar_e_ler_devolve_produto_esperado(tmp_path: Path) -> None:
    destino = _gerar(tmp_path)

    produtos = LeitorPlanilhaEntradaOpenpyxl().ler(destino)

    assert len(produtos) == 1
    produto = produtos[0]
    assert len(produto.variacoes) == 2
    assert produto.cores == ("Beige",)
    assert produto.tamanhos == ("P", "G")
