from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from loja_integrada_cadastro.infra.escritor_planilha_saida_openpyxl import (
    EscritorPlanilhaSaidaOpenpyxl,
)
from loja_integrada_cadastro.models.exceptions.erro_planilha_saida import ErroPlanilhaSaida
from loja_integrada_cadastro.models.layout_planilha_loja_integrada import (
    COLUNAS_PLANILHA_SAIDA,
)
from loja_integrada_cadastro.models.linha_planilha import LinhaPlanilha

_LINHA_PAI = LinhaPlanilha(
    tipo="com-variacao", valores={"tipo": "com-variacao", "sku": "3254002", "ativo": "S"}
)
_LINHA_FILHA = LinhaPlanilha(
    tipo="variacao",
    valores={
        "tipo": "variacao",
        "sku-pai": "3254002",
        "sku": "3254002-beige-p",
        "preco-cheio": 119.9,
        "estoque-quantidade": 10,
    },
)


def test_escreve_cabecalho_e_linhas_em_uma_aba_sheet1(tmp_path: Path) -> None:
    destino = tmp_path / "saida.xlsx"
    escritor = EscritorPlanilhaSaidaOpenpyxl()

    escritor.escrever([_LINHA_PAI, _LINHA_FILHA], destino)

    workbook = load_workbook(destino)
    assert workbook.sheetnames == ["Sheet1"]
    aba = workbook["Sheet1"]

    cabecalho = next(aba.iter_rows(min_row=1, max_row=1, values_only=True))
    assert cabecalho == COLUNAS_PLANILHA_SAIDA

    linhas = list(aba.iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 2

    posicao_sku = COLUNAS_PLANILHA_SAIDA.index("sku")
    assert linhas[0][posicao_sku] == "3254002"
    assert linhas[1][posicao_sku] == "3254002-beige-p"


def test_celulas_numericas_ficam_como_numero(tmp_path: Path) -> None:
    destino = tmp_path / "saida.xlsx"
    EscritorPlanilhaSaidaOpenpyxl().escrever([_LINHA_FILHA], destino)

    workbook = load_workbook(destino)
    aba = workbook["Sheet1"]
    linha = next(aba.iter_rows(min_row=2, max_row=2, values_only=True))

    posicao_preco = COLUNAS_PLANILHA_SAIDA.index("preco-cheio")
    posicao_estoque = COLUNAS_PLANILHA_SAIDA.index("estoque-quantidade")
    assert linha[posicao_preco] == pytest.approx(119.9)
    assert linha[posicao_estoque] == 10


def test_falha_acima_do_limite_de_linhas(tmp_path: Path) -> None:
    destino = tmp_path / "saida.xlsx"
    linhas = [_LINHA_PAI] * 9_998

    with pytest.raises(ErroPlanilhaSaida):
        EscritorPlanilhaSaidaOpenpyxl().escrever(linhas, destino)

    assert not destino.exists()
